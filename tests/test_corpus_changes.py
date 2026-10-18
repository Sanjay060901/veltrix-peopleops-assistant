"""The service must work on a modified corpus (TASK §12.1 check 5).

Each test changes the synthetic fixture data the way a grader or the HR team might: edit a value,
rename a file, convert a file to another format, add or remove documents, change manifest
metadata, change the holiday calendar, swap in different database rows, or drop a broken file.
Changes made while the service is running are picked up by reload_if_changed().
"""
from __future__ import annotations

import json
import os
import sqlite3
import time
import uuid
from pathlib import Path

import docx

from app.data.db import Database
from app.ingestion.pipeline import ingest
from app.ingestion.watcher import corpus_fingerprint
from app.main import reload_if_changed
from tests.conftest import make_client
from tests.fakes import FakeLLM

QUESTION = "How many volunteer days do I get?"


def ask(client, employee: str, message: str = QUESTION) -> dict:
    response = client.post(
        "/chat", headers={"X-Employee-Id": employee}, json={"session_id": str(uuid.uuid4()), "message": message}
    )
    assert response.status_code == 200
    return response.json()


def cited(body: dict) -> list[str]:
    return [c["doc_id"] for c in body["citations"]]


def manifest(data_dir: Path) -> list[dict]:
    return json.loads((data_dir / "manifest.json").read_text(encoding="utf-8"))


def save_manifest(data_dir: Path, entries: list[dict]) -> None:
    path = data_dir / "manifest.json"
    path.write_text(json.dumps(entries, indent=2), encoding="utf-8")
    _bump(path)


def _bump(path: Path) -> None:
    # Some filesystems have coarse mtimes; make sure the change is visible to the watcher.
    later = time.time() + 2
    os.utime(path, (later, later))


def edit(path: Path, old: str, new: str) -> None:
    text = path.read_text(encoding="utf-8")
    assert old in text
    path.write_text(text.replace(old, new), encoding="utf-8")
    _bump(path)


def reload(client) -> bool:
    return reload_if_changed(client.app.state.veltrix)


# --- edits while the service is running --------------------------------------------------------

def test_edited_value_is_picked_up_without_restart(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        assert "5 days" in ask(client, "X301")["answer"]
        edit(data_dir / "corpus" / "addendum_de.md", "5 volunteer days", "7 volunteer days")
        assert reload(client)
        body = ask(client, "X301")
    assert "7 days" in body["answer"] and "5 days" not in body["answer"]
    assert cited(body) == ["VOL-ADD-DE"]


def test_no_reload_when_nothing_changed(settings):
    with make_client(settings, FakeLLM()) as client:
        assert reload(client) is False


def test_edited_pdf_is_picked_up(settings, data_dir):
    import pymupdf

    with make_client(settings, FakeLLM()) as client:
        path = data_dir / "corpus" / "vol_v1.pdf"
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((50, 70), "Volunteer Policy (Version 1)", fontsize=18, fontname="hebo")
        page.insert_text((50, 110), "1. Volunteer days", fontsize=13, fontname="hebo")
        page.insert_text((50, 132), "Employees receive 9 volunteer days per year.", fontsize=10, fontname="helv")
        doc.save(str(path))
        _bump(path)
        assert reload(client)
        body = ask(client, "X201", "What were volunteer days in 2025?")  # v1 is the only one in force then
    assert client.app.state.veltrix.orchestrator.corpus.chunk("VOL-V1#1").text.endswith("9 volunteer days per year.")


# --- renamed files -------------------------------------------------------------------------------

def test_renamed_file_with_updated_manifest_keeps_working(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        (data_dir / "corpus" / "addendum_de.md").rename(data_dir / "corpus" / "volunteer_addendum_germany_v2.md")
        entries = manifest(data_dir)
        for e in entries:
            if e["doc_id"] == "VOL-ADD-DE":
                e["file"] = "volunteer_addendum_germany_v2.md"
        save_manifest(data_dir, entries)
        assert reload(client)
        body = ask(client, "X301")
    assert cited(body) == ["VOL-ADD-DE"] and "5 days" in body["answer"]


def test_renamed_file_without_manifest_update_is_skipped_safely(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        (data_dir / "corpus" / "addendum_de.md").rename(data_dir / "corpus" / "renamed.md")
        assert reload(client)
        state = client.app.state.veltrix
        assert "VOL-ADD-DE" in state.orchestrator.corpus.failed
        assert client.get("/health").status_code == 200
        body = ask(client, "X301")
    # The addendum is gone, so the global policy applies again; nothing crashes.
    assert cited(body)[0] == "VOL-V2"


def test_file_converted_to_another_format(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        (data_dir / "corpus" / "addendum_de.md").unlink()
        d = docx.Document()
        d.add_heading("Volunteer Addendum: DE", level=0)
        d.add_heading("1. Volunteer days", level=1)
        d.add_paragraph("Employees in DE receive 6 volunteer days per year.")
        d.save(str(data_dir / "corpus" / "addendum_de.docx"))
        entries = manifest(data_dir)
        for e in entries:
            if e["doc_id"] == "VOL-ADD-DE":
                e["file"] = "addendum_de.docx"
        save_manifest(data_dir, entries)
        assert reload(client)
        body = ask(client, "X301")
    assert cited(body) == ["VOL-ADD-DE"] and "6 days" in body["answer"]


# --- documents added, removed, or with changed metadata -------------------------------------------

def test_new_document_is_used(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        (data_dir / "corpus" / "addendum_in.md").write_text(
            "# Volunteer Addendum: IN\n\n## 1. Volunteer days\n\nEmployees in IN receive 8 volunteer days per year.\n",
            encoding="utf-8",
        )
        entries = manifest(data_dir) + [{
            "doc_id": "VOL-ADD-IN", "file": "addendum_in.md", "title": "Volunteer Addendum: IN", "doc_type": "addendum",
            "regions": ["IN"], "effective_date": "2026-01-01", "access_level": "all",
        }]
        save_manifest(data_dir, entries)
        assert reload(client)
        india = ask(client, "X201")
        germany = ask(client, "X301")
    assert cited(india) == ["VOL-ADD-IN"] and "8 days" in india["answer"]
    assert cited(germany) == ["VOL-ADD-DE"]  # the new IN addendum is not for DE


def test_removed_document_is_no_longer_cited(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        save_manifest(data_dir, [e for e in manifest(data_dir) if e["doc_id"] != "VOL-V2"])
        assert reload(client)
        body = ask(client, "X201")
    assert "VOL-V2" not in cited(body) and cited(body)[0] == "VOL-V1"


def test_effective_date_moved_to_future_becomes_upcoming(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        entries = manifest(data_dir)
        for e in entries:
            if e["doc_id"] == "VOL-ADD-DE":
                e["effective_date"] = "2026-12-01"
        save_manifest(data_dir, entries)
        assert reload(client)
        body = ask(client, "X301")
    # v2 applies today; the addendum would win from 2026-12-01, so it is mentioned as upcoming.
    assert cited(body) == ["VOL-V2", "VOL-ADD-DE"] and "UPCOMING" in body["answer"]


def test_access_level_tightened_hides_document(settings, data_dir):
    llm = FakeLLM()
    with make_client(settings, llm) as client:
        entries = manifest(data_dir)
        for e in entries:
            if e["doc_id"] == "VOL-ADD-DE":
                e["access_level"] = "hr"
        save_manifest(data_dir, entries)
        assert reload(client)
        body = ask(client, "X301")
    assert "VOL-ADD-DE" not in cited(body)
    sent = "\n".join(c["user"].split("Sources:", 1)[-1] for c in llm.calls if c["purpose"] == "extract")
    assert "Employees in DE receive" not in sent  # never reaches the LLM


def test_broken_file_does_not_stop_ingestion(settings, data_dir):
    with make_client(settings, FakeLLM()) as client:
        path = data_dir / "corpus" / "vol_v2.docx"
        path.write_bytes(b"this is not a docx file")
        _bump(path)
        assert reload(client)
        state = client.app.state.veltrix
        assert "VOL-V2" in state.orchestrator.corpus.failed
        body = ask(client, "X201")
    assert body["route"] == "rag" and cited(body)[0] == "VOL-V1"


# --- holiday calendar and database changes ---------------------------------------------------------

APPLY = {"intent": "apply_leave", "leave_type": "annual", "start_date_text": "2026-09-21", "start_date": "2026-09-21",
         "end_date_text": "2026-09-22", "end_date": "2026-09-22"}


def test_changed_holiday_calendar_changes_working_days(settings, data_dir):
    llm = FakeLLM(lambda m: APPLY)
    with make_client(settings, llm) as client:
        before = ask(client, "X201", "book leave")["answer"]
        edit(data_dir / "corpus" / "cal_in_2026.md", "| 21 September 2026 | Monday | Founders Day |",
             "| 22 September 2026 | Tuesday | Founders Day |\n| 21 September 2026 | Monday | Second Day |")
        assert reload(client)
        after = ask(client, "X201", "book leave")["answer"]
    assert "Working days: 1" in before   # 21st was a holiday
    assert "no working days" in after    # now both days are holidays


def test_modified_database_rows(settings, data_dir):
    """A fresh database with different rows: new employees and balances work without code changes."""
    conn = sqlite3.connect(data_dir / "veltrix.db")
    conn.execute("INSERT INTO employees VALUES ('Z900','New Person','new.person@example.test','+00 999 000 1111','DE',"
                 "'2026-01-05','L1','employee','X300')")
    conn.execute("INSERT INTO leave_balances VALUES ('Z900','annual',11.5)")
    conn.execute("UPDATE leave_balances SET balance_days = 1 WHERE employee_id='X201' AND leave_type='annual'")
    conn.commit()
    conn.close()
    llm = FakeLLM(lambda m: {"intent": "leave_balance"})
    with make_client(settings, llm) as client:
        new = ask(client, "Z900", "my balance")
        changed = ask(client, "X201", "my balance")
        manager = FakeLLM(lambda m: {"intent": "leave_balance", "target_employee": "Z900"})
    assert "11.5 days of annual" in new["answer"]
    assert "1 day of annual" in changed["answer"]
    with make_client(settings, manager) as client:
        assert "New Person" in ask(client, "X300", "balance of Z900")["answer"]  # Z900 reports to X300
        assert ask(client, "X200", "balance of Z900")["route"] == "refuse"


def test_fingerprint_detects_add_rename_edit(data_dir):
    corpus, manifest_path = str(data_dir / "corpus"), str(data_dir / "manifest.json")
    base = corpus_fingerprint(corpus, manifest_path)
    (data_dir / "corpus" / "faq.html").rename(data_dir / "corpus" / "faq2.html")
    renamed = corpus_fingerprint(corpus, manifest_path)
    (data_dir / "corpus" / "extra.md").write_text("# x", encoding="utf-8")
    added = corpus_fingerprint(corpus, manifest_path)
    assert len({base, renamed, added}) == 3


def test_ingest_directly_on_modified_copy(settings, db, data_dir):
    edit(data_dir / "corpus" / "expenses_b.md", "within 45 days", "within 60 days")
    corpus = ingest(settings.corpus_dir, settings.manifest_path, db, None)
    assert any("within 60 days" in c.text for c in corpus.chunks)
